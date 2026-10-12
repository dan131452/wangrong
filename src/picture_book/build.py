"""把故事文案与插画排版成绘本：PDF（reportlab）+ PPTX（python-pptx）。

用法：
    uv run picture-book check            # 检查插画就位情况
    uv run picture-book build            # 生成 dist/ 下的 PDF 与 PPTX（缺图时自动用占位图）
    uv run picture-book build --only pdf # 只重新生成其中一种
    uv run picture-book gen              # 调用生图 API 批量生成插画（配置见 .env）
"""

from __future__ import annotations

import argparse
import sys
from math import ceil
from pathlib import Path

from .fonts import first_existing
from .placeholders import placeholder_for
from .story import (
    CACHE_DIR,
    DIST_DIR,
    Book,
    Page,
    available_book_ids,
    find_image,
    load_book,
    load_books,
)

# 调色板（0-255 RGB，PDF/PPTX 共用）
CREAM = (254, 248, 236)     # 卡片底色
INK = (107, 63, 42)         # 暖棕：标题
BODY = (74, 58, 46)         # 深棕：正文
PAGE_NUM = (154, 143, 130)  # 页码灰
ACCENT = (199, 124, 58)     # 暖橙：标题强调条
SHADOW = (150, 132, 108)    # 卡片投影

# 中文排版禁则：这些标点不能出现在行首
NO_LINE_START = set("，。！？；：、）】》〉”’…—％‰,.!?;:%)]}")

# 版式常量
PDF_CARD_W = 396.0      # 故事页文字卡片宽（pt）
PDF_CARD_PAD = 24.0     # 卡片内边距
PDF_FONT_SIZE = 13.5
PDF_LEADING = 22.0
PDF_FRAME_OUT = 14.0    # 页面装饰外框 inset（pt）
PDF_FRAME_IN = 20.0     # 页面装饰内框 inset（pt）
PPTX_W, PPTX_H = 13.333, 7.5  # 英寸


# ---------------------------------------------------------------- 图像工具

def cover_crop(src: Path, out_path: Path, ratio: float) -> Path:
    """居中裁剪成目标宽高比（cover 模式），输出 JPEG。"""
    from PIL import Image, ImageOps

    out_path.parent.mkdir(parents=True, exist_ok=True)
    img = ImageOps.exif_transpose(Image.open(src)).convert("RGB")
    w, h = img.size
    if w / h > ratio:  # 太宽，裁两侧
        new_w = int(h * ratio)
        box = ((w - new_w) // 2, 0, (w - new_w) // 2 + new_w, h)
    else:  # 太高，裁上下
        new_h = int(w / ratio)
        box = (0, (h - new_h) // 2, w, (h - new_h) // 2 + new_h)
    img.crop(box).save(out_path, "JPEG", quality=90)
    return out_path


def resolve_image(page: Page, book_id: str) -> tuple[Path, bool]:
    """返回 (原始图片路径, 是否为占位图)，由各构建器自行裁剪。"""
    real = find_image(page)
    if real is not None:
        return real, False
    return placeholder_for(book_id, page.no, page.title), True


def resolve_images(book: Book) -> tuple[dict[int, Path], list[int]]:
    """一次解析整本书的图，返回 ({页码: 路径}, 用了占位图的页码列表)。"""
    images: dict[int, Path] = {}
    placeholders: list[int] = []
    for page in book.pages:
        path, is_ph = resolve_image(page, book.id)
        images[page.no] = path
        if is_ph:
            placeholders.append(page.no)
    return images, placeholders


def wrap_text(text: str, font_name: str, size: float, max_width: float) -> list[str]:
    """按像素宽度逐字换行（中文场景够用），禁则标点悬挂在上一行行尾。"""
    from reportlab.pdfbase import pdfmetrics

    lines: list[str] = []
    line = ""
    for ch in text:
        if pdfmetrics.stringWidth(line + ch, font_name, size) <= max_width:
            line += ch
        elif ch in NO_LINE_START and line:
            line += ch  # 允许轻微超宽，把标点留在行尾
            lines.append(line)
            line = ""
        else:
            lines.append(line)
            line = ch
    if line:
        lines.append(line)
    return lines


# ---------------------------------------------------------------- PDF

def _register_pdf_fonts() -> dict[str, str]:
    """注册中文字体，返回 {"regular": 字体名, "bold": 字体名}；全缺时用内置 CID 字体兜底。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFError, TTFont

    names: dict[str, str] = {}
    for slot in ("regular", "bold"):
        path = first_existing(slot)
        if path is not None:
            try:
                name = path.stem.upper()
                pdfmetrics.registerFont(TTFont(name, str(path), subfontIndex=0))
                names[slot] = name
                continue
            except (TTFError, OSError) as exc:
                print(f"提示：字体 {path.name} 注册失败（{exc}），改用内置 CID 字体。", file=sys.stderr)
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        names[slot] = "STSong-Light"
    return names


def build_pdf(book: Book, images: dict[int, Path], out_path: Path) -> None:
    from reportlab.lib.colors import Color
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    fonts = _register_pdf_fonts()
    PW, PH = landscape(A4)  # 842 x 595 pt

    def rgb(c: tuple[int, int, int], alpha: float = 1.0) -> Color:
        return Color(c[0] / 255, c[1] / 255, c[2] / 255, alpha=alpha)

    def page_frame() -> None:
        """页面四周的双线装饰边框：外细棕线 + 内奶油粗线。"""
        c.setLineWidth(1.2)
        c.setStrokeColor(rgb(INK, alpha=0.35))
        c.roundRect(PDF_FRAME_OUT, PDF_FRAME_OUT, PW - 2 * PDF_FRAME_OUT, PH - 2 * PDF_FRAME_OUT, 10, stroke=1, fill=0)
        c.setLineWidth(2.2)
        c.setStrokeColor(rgb(CREAM, alpha=0.65))
        c.roundRect(PDF_FRAME_IN, PDF_FRAME_IN, PW - 2 * PDF_FRAME_IN, PH - 2 * PDF_FRAME_IN, 8, stroke=1, fill=0)

    c = canvas.Canvas(str(out_path), pagesize=(PW, PH))
    c.setTitle(book.title)

    for page in book.pages:
        c.drawImage(
            str(cover_crop(images[page.no], CACHE_DIR / f"p{page.no:02d}-pdf.jpg", PW / PH)),
            0, 0, width=PW, height=PH,
        )
        page_frame()

        if page.kind == "cover":
            band_h = 132
            c.setFillColor(rgb(CREAM, alpha=0.93))
            c.rect(0, 0, PW, band_h, stroke=0, fill=1)
            c.setFillColor(rgb(INK))
            c.setFont(fonts["bold"], 34)
            c.drawCentredString(PW / 2, band_h - 62, page.title)
            c.setFillColor(rgb(PAGE_NUM))
            c.setFont(fonts["regular"], 13)
            c.drawCentredString(PW / 2, band_h - 96, book.subtitle)

        elif page.kind == "back":
            cw, ch = 500.0, 200.0
            cx, cy = (PW - cw) / 2, (PH - ch) / 2
            # 封底卡片居中，正压在插画视觉中心上，必须不透明才能保证文字干净
            c.setFillColor(rgb(SHADOW, alpha=0.25))
            c.roundRect(cx + 5, cy - 5, cw, ch, 16, stroke=0, fill=1)
            c.setFillColor(rgb(CREAM, alpha=1.0))
            c.roundRect(cx, cy, cw, ch, 16, stroke=0, fill=1)
            c.setFillColor(rgb(ACCENT))
            c.roundRect(cx + 22, cy + ch - 44, 5, 26, 2.5, stroke=0, fill=1)
            c.setFillColor(rgb(INK))
            c.setFont(fonts["bold"], 18)
            c.drawString(cx + 38, cy + ch - 38, page.title)
            c.setFillColor(rgb(ACCENT))
            c.setFont(fonts["regular"], 9)
            c.drawString(cx + 36, cy + ch - 60, "◆")
            c.setFillColor(rgb(BODY))
            c.setFont(fonts["regular"], 13.5)
            for i, line in enumerate(wrap_text(page.text or "", fonts["regular"], 13.5, cw - 72)):
                c.drawString(cx + 36, cy + ch - 86 - i * 24, line)
            c.setFillColor(rgb(PAGE_NUM))
            c.setFont(fonts["regular"], 10.5)
            c.drawCentredString(PW / 2, cy + 24, "—— ◆ 完 ◆ ——")

        else:  # story
            text_w = PDF_CARD_W - 2 * PDF_CARD_PAD
            lines = wrap_text(page.text or "", fonts["regular"], PDF_FONT_SIZE, text_w)
            ch = 20 + 32 + len(lines) * PDF_LEADING + 34
            cx, cy = 44.0, 44.0
            c.setFillColor(rgb(SHADOW, alpha=0.25))
            c.roundRect(cx + 5, cy - 5, PDF_CARD_W, ch, 16, stroke=0, fill=1)
            c.setFillColor(rgb(CREAM, alpha=0.95))
            c.roundRect(cx, cy, PDF_CARD_W, ch, 16, stroke=0, fill=1)
            c.setFillColor(rgb(ACCENT))
            c.roundRect(cx + 14, cy + ch - 46, 5, 26, 2.5, stroke=0, fill=1)
            c.setFillColor(rgb(INK))
            c.setFont(fonts["bold"], 17)
            c.drawString(cx + PDF_CARD_PAD + 8, cy + ch - 40, page.title)
            c.setFillColor(rgb(BODY))
            c.setFont(fonts["regular"], PDF_FONT_SIZE)
            for i, line in enumerate(lines):
                c.drawString(cx + PDF_CARD_PAD, cy + ch - 70 - i * PDF_LEADING, line)
            c.setFillColor(rgb(PAGE_NUM))
            c.setFont(fonts["regular"], 10)
            c.drawRightString(cx + PDF_CARD_W - 20, cy + 14, f"· {page.no} ·")

        c.showPage()

    c.save()


# ---------------------------------------------------------------- PPTX

def build_pptx(book: Book, images: dict[int, Path], out_path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    cream = RGBColor(*CREAM)
    ink = RGBColor(*INK)
    body = RGBColor(*BODY)
    page_num = RGBColor(*PAGE_NUM)
    accent = RGBColor(*ACCENT)
    shadow_fill = RGBColor(216, 204, 180)
    frame_inner = RGBColor(198, 182, 158)
    ea_font = "微软雅黑"

    def style_run(run, text: str, size: float, bold: bool, color: RGBColor) -> None:
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = ea_font
        run.font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rPr.append(rPr.makeelement(qn("a:ea"), {"typeface": ea_font}))

    def add_frame(slide) -> None:
        """页面四周双线装饰边框：外奶油粗线 + 内暖棕细线。"""
        for inset, color, width in ((0.10, cream, 2.4), (0.17, frame_inner, 1.0)):
            frame = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(inset), Inches(inset),
                prs.slide_width - Inches(2 * inset), prs.slide_height - Inches(2 * inset),
            )
            frame.fill.background()
            frame.line.color.rgb = color
            frame.line.width = Pt(width)
            frame.shadow.inherit = False

    def add_accent_bar(slide, x, y) -> None:
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(0.07), Inches(0.42))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()
        bar.shadow.inherit = False

    prs = Presentation()
    prs.slide_width = Inches(PPTX_W)
    prs.slide_height = Inches(PPTX_H)
    blank = prs.slide_layouts[6]

    for page in book.pages:
        slide = prs.slides.add_slide(blank)
        img = cover_crop(images[page.no], CACHE_DIR / f"p{page.no:02d}-pptx.jpg", PPTX_W / PPTX_H)
        slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
        add_frame(slide)

        def add_card(slide, x, y, w, h):
            sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.05), y + Inches(0.05), w, h)
            sh.fill.solid()
            sh.fill.fore_color.rgb = shadow_fill
            sh.line.fill.background()
            sh.shadow.inherit = False
            shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
            shp.fill.solid()
            shp.fill.fore_color.rgb = cream
            shp.line.fill.background()
            shp.shadow.inherit = False
            tf = shp.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Inches(0.28)
            tf.margin_top = tf.margin_bottom = Inches(0.16)
            return tf

        if page.kind == "cover":
            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, prs.slide_height - Inches(1.75), prs.slide_width, Inches(1.75))
            band.fill.solid()
            band.fill.fore_color.rgb = cream
            band.line.fill.background()
            band.shadow.inherit = False
            tf = band.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            style_run(p.add_run(), page.title, 40, True, ink)
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            style_run(p2.add_run(), book.subtitle, 15, False, page_num)

        elif page.kind == "back":
            tf = add_card(slide, (prs.slide_width - Inches(6.4)) / 2, (prs.slide_height - Inches(2.5)) / 2, Inches(6.4), Inches(2.5))
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            style_run(p.add_run(), page.title, 21, True, ink)
            p_mid = tf.add_paragraph()
            p_mid.alignment = PP_ALIGN.CENTER
            p_mid.space_before = Pt(6)
            style_run(p_mid.add_run(), "◆", 10, False, accent)
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            p2.space_before = Pt(6)
            style_run(p2.add_run(), page.text or "", 15, False, body)
            p3 = tf.add_paragraph()
            p3.alignment = PP_ALIGN.CENTER
            p3.space_before = Pt(10)
            style_run(p3.add_run(), "—— ◆ 完 ◆ ——", 11, False, page_num)

        else:  # story
            # 卡片高度按正文行数自适应，正文加长也不会溢出卡片
            body_lines = max(1, ceil(len(page.text or "") / 22))  # 14.5pt 中文在 5.0in 卡片约 22 字/行
            card_h = Inches(max(2.35, 1.05 + body_lines * 0.24))
            card_y = prs.slide_height - card_h - Inches(0.5)
            tf = add_card(slide, Inches(0.5), card_y, Inches(5.0), card_h)
            add_accent_bar(slide, Inches(0.60), card_y + Inches(0.28))
            p = tf.paragraphs[0]
            style_run(p.add_run(), page.title, 19, True, ink)
            p2 = tf.add_paragraph()
            p2.space_before = Pt(8)
            style_run(p2.add_run(), page.text or "", 14.5, False, body)

            num = slide.shapes.add_textbox(prs.slide_width - Inches(1.2), prs.slide_height - Inches(0.55), Inches(0.9), Inches(0.35))
            ntf = num.text_frame
            ntf.margin_left = ntf.margin_right = ntf.margin_top = ntf.margin_bottom = 0
            np_ = ntf.paragraphs[0]
            np_.alignment = PP_ALIGN.RIGHT
            style_run(np_.add_run(), f"· {page.no} ·", 11, False, page_num)

    prs.save(str(out_path))


# ---------------------------------------------------------------- 命令行

def display_pad(text: str, width: int) -> str:
    """按“CJK 字符算 2 格”补空格，让中英文混排的列表对齐。"""
    used = sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
    return text + " " * max(0, width - used)


def cmd_prompts() -> int:
    """从 data/books/*.json 生成 data/prompts.md，保证提示词清单与数据源一致。"""
    parts = [
        "# 出图提示词清单，由 data/books/*.json 自动生成",
        "",
        "> 由 `uv run picture-book prompts` 重新生成；改提示词请改 books 里的 JSON，不要直接改本文件。",
        "> 比例选 **横版**（API 生图用 1536x1024；手动工具选 4:3 或 3:2 均可，管道会自动裁剪）。",
        "> 每页建议生成 2~4 张挑最好的一张，按括号里的文件名保存到 `assets/images/`。",
        "> 角色一致性：豆包把图发进对话当参考（“参考这张图里的角色，……”）；即梦用“智能参考”。",
        "> 两本书共用诸葛亮形象：先出《草船借箭》P3 的诸葛亮，同一形象贯穿两本书。",
        "",
        "**推荐工具（2026-09 实测现状）：**",
        "1. **豆包**（app 或 doubao.com）：Seedream 生图免费，量大管饱，参考图一致性最好 —— 首选",
        "2. **通义万相**（wanx.aliyun.com）：新用户送 100 张生成额度",
        "3. **即梦**（jimeng.jianying.com）：每日免费积分已缩水（约 3~6 张/天），适合补图",
        "4. **本项目 gen 命令**：`uv run picture-book gen`（走 .env 里配置的 API）",
        "",
    ]
    for book in load_books():
        parts.append("---\n")
        parts.append(f"# 《{book.title}》\n")
        for p in book.pages:
            label = "封面" if p.kind == "cover" else ("封底" if p.kind == "back" else "故事")
            parts.append(f"**P{p.no} {p.title}**（{label}，`{p.image}`）\n")
            parts.append(f"{p.prompt}\n")
    DATA_DIR_PATH = DIST_DIR.parent / "data"
    (DATA_DIR_PATH / "prompts.md").write_text("\n".join(parts), encoding="utf-8")
    print("已生成 data/prompts.md")
    return 0


def cmd_check() -> int:
    missing_total = 0
    for book in load_books():
        missing = [p for p in book.pages if find_image(p) is None]
        missing_total += len(missing)
        print(f"《{book.title}》（{book.id}）共 {len(book.pages)} 页")
        for page in book.pages:
            ok = find_image(page) is not None
            mark = "✓" if ok else "✗ 缺"
            print(f"  [{mark}] P{page.no:02d} {display_pad(page.title, 14)}期望文件: assets/images/{page.image}")
        print(f"  -> 就绪 {len(book.pages) - len(missing)}/{len(book.pages)}\n")
    if missing_total:
        print("请用 data/prompts.md 中的提示词出图，或 uv run picture-book gen 自动生成。")
        return 1
    print("全部插画就绪，运行 uv run picture-book build 即可出成品。")
    return 0


def cmd_build(book_id: str | None = None, only: str | None = None) -> int:
    books = [load_book(book_id)] if book_id else load_books()
    DIST_DIR.mkdir(parents=True, exist_ok=True)

    images_by_book: dict[str, dict[int, Path]] = {}
    any_placeholder = False
    outputs: list[str] = []
    for book in books:
        images, placeholders = resolve_images(book)
        images_by_book[book.id] = images
        any_placeholder = any_placeholder or bool(placeholders)
        if only in (None, "pdf"):
            pdf_path = DIST_DIR / f"{book.title}.pdf"
            build_pdf(book, images, pdf_path)
            outputs.append(f"PDF : {pdf_path}")
        if only in (None, "pptx"):
            pptx_path = DIST_DIR / f"{book.title}.pptx"
            build_pptx(book, images, pptx_path)
            outputs.append(f"PPTX: {pptx_path}")
    if only in (None, "web"):
        from .web import build_web

        build_web(books, images_by_book, DIST_DIR / "绘本阅读器.html")

    print("\n".join(outputs))
    if any_placeholder:
        print("\n注意：有页面使用了占位图（见上方 check）。")
        print("出图后放进 assets/images/ 再重新 build 即可。")
    else:
        print("\n全部插画已就位，成品可直接使用。")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(prog="picture-book", description="AIGC 数字绘本排版构建工具")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("check", help="检查所有绘本插画就位情况")
    sub.add_parser("prompts", help="重新生成 data/prompts.md")
    build_parser = sub.add_parser("build", help="生成 PDF、PPTX 与网页阅读器")
    build_parser.add_argument("--book", choices=available_book_ids(), help="只构建指定绘本")
    build_parser.add_argument("--only", choices=("pdf", "pptx", "web"), help="只生成其中一种成品")
    gen_parser = sub.add_parser("gen", help="调用生图 API 批量生成插画（配置见 .env）")
    gen_parser.add_argument("--book", choices=available_book_ids(), help="只生成指定绘本的插画")
    gen_parser.add_argument("--page", type=int, help="只生成指定页码")
    gen_parser.add_argument("--force", action="store_true", help="已有图片也重新生成")
    serve_parser = sub.add_parser("serve", help="启动本地 AIGC 生成器网页（选故事→点生成→阅读）")
    serve_parser.add_argument("--port", type=int, default=8765)
    serve_parser.add_argument("--demo", action="store_true", help="演示模式：不调真实 API，用占位图模拟生成")
    args = parser.parse_args()

    try:
        if args.command == "check":
            return cmd_check()
        if args.command == "prompts":
            return cmd_prompts()
        if args.command == "gen":
            from .gen import cmd_gen

            return cmd_gen(book_id=args.book, page=args.page, force=args.force)
        if args.command == "serve":
            from .server import cmd_serve

            return cmd_serve(port=args.port, demo=args.demo)
        return cmd_build(book_id=args.book, only=args.only)
    except (FileNotFoundError, ValueError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
